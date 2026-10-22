"""Детали по пикам из ручной работы Тарабрина.

Для каждого пика F14/F16 выводит:
* пики исходного спектра, вызывающие source_overlap, и их интенсивности
  относительно исходного пика/пика дериватизата;
* все другие серии, делящие пик дериватизата (conflict_overlap).
"""

from __future__ import annotations

import argparse
import ast
import bisect
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document

from src.core.spectrum import load_spectrum

DOCX = Path(
    "C:/Users/mvs/Work/Лаборатория/Аналитика/Тарабрин аналитика альфа-версия.docx"
)
SETS = {
    "Fulvagra F14": {
        "src": "data/real_sets/Fulvagra F14/Fulvagra_prep HPLC_Fr14_ESI_1e6_480k_128_m-s_averaged.json",
        "dmet": "data/real_sets/Fulvagra F14/Fulvagra_prep HPLC_FR14_CD3OD mod_ESI_1e6_480k_128_m-s_averaged.json",
        "dacet": "data/real_sets/Fulvagra F14/Fulvagra_prep HPLC_FR14_AC-D4_ESI_1e6_480k_128_m-s_averaged.json",
        "mass_table": 0,
        "count_table": 2,
    },
    "Fulvagra F16": {
        "src": "data/real_sets/Fulvagra F16/Fulvagra_prep HPLC_Fr16_ESI_1e6_480k_128_m-s_averaged.json",
        "dmet": "data/real_sets/Fulvagra F16/Fulvagra_prep HPLC_FR16_CD3OD_ESI_1e6_480k_128_m-s__averaged.json",
        "dacet": "data/real_sets/Fulvagra F16/Fulvagra_prep HPLC_FR16_AC-D4_ESI_1e6_480k_128_m-s_averaged.json",
        "mass_table": 1,
        "count_table": 3,
    },
}
PPM_LOOKUP = 5.0


def _clean(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split())


def _int(value: str) -> int | None:
    m = re.match(r"\s*(\d+)", _clean(value))
    return int(m.group(1)) if m else None


def _manual(doc: Document, cfg: dict) -> list[dict]:
    counts: dict[str, dict] = {}
    for row in doc.tables[cfg["count_table"]].rows[2:]:
        c = [_clean(x.text) for x in row.cells]
        if len(c) >= 4 and c[0]:
            counts[c[0]] = {"formula": c[1], "cooh": _int(c[2]), "oh": _int(c[3])}
    out = []
    for row in doc.tables[cfg["mass_table"]].rows[2:]:
        c = [_clean(x.text) for x in row.cells]
        if len(c) < 3 or c[0] not in counts:
            continue
        out.append({"number": c[0], "mass": float(c[1]), **counts[c[0]]})
    return out


def _details(value) -> list[dict]:
    if isinstance(value, list):
        return [x for x in value if isinstance(x, dict)]
    if isinstance(value, str) and value.strip():
        try:
            parsed = ast.literal_eval(value)
            return (
                [x for x in parsed if isinstance(x, dict)]
                if isinstance(parsed, list)
                else []
            )
        except (ValueError, SyntaxError):
            return []
    return []


def _nearest(
    masses: np.ndarray, intensities: np.ndarray, target: float, ppm=PPM_LOOKUP
):
    tol = abs(target) * ppm * 1e-6
    lo = int(np.searchsorted(masses, target - tol, side="left"))
    hi = int(np.searchsorted(masses, target + tol, side="right"))
    if hi <= lo:
        return None, None
    idx = lo + int(np.argmax(intensities[lo:hi]))
    return float(masses[idx]), float(intensities[idx])


class SeriesIndex:
    """Все ступени всех серий: m/z → (row, step)."""

    def __init__(self, df: pd.DataFrame, detail_col: str, tolerance=0.001):
        self.tolerance = tolerance
        self.df = df
        self.items = []
        for row_idx, value in enumerate(df[detail_col]):
            for detail in _details(value):
                mz = detail.get("mz")
                if mz is not None:
                    self.items.append(
                        (float(mz), int(row_idx), int(detail.get("n", 0)))
                    )
        self.items.sort(key=lambda x: x[0])
        self.masses = [x[0] for x in self.items]

    def query(self, mz: float) -> list[tuple[int, int, float]]:
        lo = bisect.bisect_left(self.masses, mz - self.tolerance)
        hi = bisect.bisect_right(self.masses, mz + self.tolerance)
        return [(row, step, m) for m, row, step in self.items[lo:hi]]


def _run(docx: Path, results: Path, out: Path, datasets: str = "all") -> None:
    doc = Document(docx)
    result_rows = []
    summaries = {}
    wanted = {x.strip().lower() for x in datasets.split(",")}
    selected = [
        name
        for name in SETS
        if datasets.strip().lower() in ("all", "real") or name.lower() in wanted
    ]
    for dataset in selected:
        cfg = SETS[dataset]
        result_path = results / dataset / "result_overlap_only.csv.gz"
        if not result_path.is_file():
            result_path = results / dataset / "result.csv.gz"
        df = pd.read_csv(result_path, low_memory=False)
        src = load_spectrum(cfg["src"])
        dmet = load_spectrum(cfg["dmet"])
        dacet = load_spectrum(cfg["dacet"])
        spec_by_deriv = {"CD3": dmet, "CD3CO": dacet}
        details_by_deriv = {"CD3": "N_COOH_details", "CD3CO": "N_OH_details"}
        index_by_deriv = {
            name: SeriesIndex(df, details_by_deriv[name]) for name in ("CD3", "CD3CO")
        }
        src_mass = np.asarray(src.table["mass"], dtype=float)
        src_int = np.asarray(src.table["intensity"], dtype=float)
        order = np.argsort(src_mass)
        src_mass, src_int = src_mass[order], src_int[order]
        mass_cols = {"CD3": "N_COOH", "CD3CO": "N_OH"}
        raw_cols = {"CD3": "N_COOH_raw", "CD3CO": "N_OH_raw"}

        summaries[dataset] = {"peaks": []}
        for manual in _manual(doc, cfg):
            mass = manual["mass"]
            idx = (df["mass"].astype(float) - mass).abs().idxmin()
            row = df.loc[idx]
            parent_mass, parent_int = _nearest(src_mass, src_int, mass)
            peak_summary = {
                "number": manual["number"],
                "mass": mass,
                "manual_formula": manual["formula"],
                "chosen_formula": str(row["brutto"]),
                "manual_cooh": manual["cooh"],
                "manual_oh": manual["oh"],
                "software_cooh": int(row["N_COOH"]),
                "software_oh": int(row["N_OH"]),
                "software_cooh_raw": int(row["N_COOH_raw"]),
                "software_oh_raw": int(row["N_OH_raw"]),
                "source_events": 0,
                "conflict_events": 0,
                "conflict_series": 0,
            }
            for deriv, details_col in details_by_deriv.items():
                spec = spec_by_deriv[deriv]
                dm = np.asarray(spec.table["mass"], dtype=float)
                di = np.asarray(spec.table["intensity"], dtype=float)
                dorder = np.argsort(dm)
                dm, di = dm[dorder], di[dorder]
                details = _details(row[details_col])
                source_seen: set[tuple] = set()
                for detail in details:
                    step = int(detail.get("n", 0))
                    step_mz = float(detail["mz"])
                    _, deriv_int = _nearest(dm, di, step_mz)
                    for source_mz in detail.get("src_mz") or []:
                        smz, sint = _nearest(src_mass, src_int, float(source_mz))
                        if sint is None:
                            continue
                        source_key = (deriv, step, round(float(smz), 4))
                        if source_key in source_seen:
                            continue
                        source_seen.add(source_key)
                        result_rows.append(
                            {
                                "dataset": dataset,
                                "number": manual["number"],
                                "manual_mass": mass,
                                "manual_formula": manual["formula"],
                                "chosen_formula": str(row["brutto"]),
                                "derivative": deriv,
                                "step": step,
                                "step_mz": step_mz,
                                "event_type": "source_overlap",
                                "source_mz": smz,
                                "source_intensity": sint,
                                "parent_intensity": parent_int,
                                "source_to_parent": (
                                    sint / parent_int if parent_int else None
                                ),
                                "deriv_intensity": deriv_int,
                                "source_to_deriv": (
                                    sint / deriv_int if deriv_int else None
                                ),
                                "other_mass": None,
                                "other_formula": None,
                                "other_step": None,
                                "other_raw": None,
                                "other_confirmed": None,
                            }
                        )
                        peak_summary["source_events"] += 1
                    if detail.get("conflict_overlap"):
                        others = index_by_deriv[deriv].query(step_mz)
                        seen = set()
                        for other_row, other_step, _other_mz in others:
                            if other_row == idx:
                                continue
                            other = df.iloc[other_row]
                            key = (
                                float(other["mass"]),
                                str(other["brutto"]),
                                int(other_step),
                            )
                            if key in seen:
                                continue
                            seen.add(key)
                            result_rows.append(
                                {
                                    "dataset": dataset,
                                    "number": manual["number"],
                                    "manual_mass": mass,
                                    "manual_formula": manual["formula"],
                                    "chosen_formula": str(row["brutto"]),
                                    "derivative": deriv,
                                    "step": step,
                                    "step_mz": step_mz,
                                    "event_type": "conflict_overlap",
                                    "source_mz": None,
                                    "source_intensity": None,
                                    "parent_intensity": parent_int,
                                    "source_to_parent": None,
                                    "deriv_intensity": deriv_int,
                                    "source_to_deriv": None,
                                    "other_mass": float(other["mass"]),
                                    "other_formula": str(other["brutto"]),
                                    "other_step": int(other_step),
                                    "other_raw": int(other.get(raw_cols[deriv], 0)),
                                    "other_confirmed": int(
                                        other.get(mass_cols[deriv], 0)
                                    ),
                                }
                            )
                            peak_summary["conflict_events"] += 1
                        peak_summary["conflict_series"] += len(seen)
            summaries[dataset]["peaks"].append(peak_summary)

    out.mkdir(parents=True, exist_ok=True)
    df_out = pd.DataFrame(result_rows)
    df_out.to_csv(out / "manual_peak_overlap_details.csv", index=False)
    (out / "manual_peak_overlap_details.json").write_text(
        json.dumps(summaries, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # Markdown: группируем по пику, чтобы всё было видно без CSV.
    lines = [
        "# Детали source/conflict overlap по пикам Тарабрина",
        "",
        "`I_source/I_parent` — интенсивность source-пика, попавшего в окно шага,",
        "делённая на интенсивность исходного пика из ручной таблицы.",
        "`I_source/I_deriv` — то же относительно пика дериватизата на этом шаге.",
        "В conflict-таблицах перечислены все другие серии, у которых есть ступень",
        "в пределах ±0.001 m/z от того же пика дериватизата.",
        "",
    ]
    for dataset, payload in summaries.items():
        lines += [f"## {dataset}", ""]
        for p in payload["peaks"]:
            lines += [
                f"### {p['number']}  m/z {p['mass']:.5f}  "
                f"{p['manual_formula']} → {p['chosen_formula']}",
                "",
                f"COOH ручн./ПО: {p['manual_cooh']} / {p['software_cooh']} "
                f"(raw {p['software_cooh_raw']}); "
                f"OH ручн./ПО: {p['manual_oh']} / {p['software_oh']} "
                f"(raw {p['software_oh_raw']}).",
                "",
            ]
            sub = df_out[
                (df_out["dataset"] == dataset) & (df_out["number"] == p["number"])
            ]
            src_sub = sub[sub["event_type"] == "source_overlap"]
            con_sub = sub[sub["event_type"] == "conflict_overlap"]
            lines += ["**Source overlap**", ""]
            if src_sub.empty:
                lines += ["нет", ""]
            else:
                lines += [
                    "| derivative | step | step m/z | source m/z | I_source | I_parent | I_source/I_parent | I_deriv |",
                    "|---|---:|---:|---:|---:|---:|---:|---:|",
                ]
                for _, r in src_sub.iterrows():
                    lines.append(
                        f"| {r['derivative']} | {r['step']} | {r['step_mz']:.5f} | "
                        f"{r['source_mz']:.5f} | {r['source_intensity']:.3g} | "
                        f"{r['parent_intensity']:.3g} | {r['source_to_parent']:.3g} | "
                        f"{r['deriv_intensity']:.3g} |"
                        if pd.notna(r["source_to_parent"])
                        and pd.notna(r["deriv_intensity"])
                        else f"| {r['derivative']} | {r['step']} | {r['step_mz']:.5f} | "
                        f"{r['source_mz']:.5f} | {r['source_intensity']:.3g} | — | — | — |"
                    )
                lines.append("")
            lines += ["**Conflict overlap: другие серии**", ""]
            if con_sub.empty:
                lines += ["нет", ""]
            else:
                lines += [
                    "| derivative | step | step m/z | other m/z | other formula | "
                    "other step | other raw | other confirmed |",
                    "|---|---:|---:|---:|---|---:|---:|---:|",
                ]
                for _, r in con_sub.iterrows():
                    lines.append(
                        f"| {r['derivative']} | {r['step']} | {r['step_mz']:.5f} | "
                        f"{r['other_mass']:.5f} | {r['other_formula']} | "
                        f"{int(r['other_step'])} | {int(r['other_raw'])} | "
                        f"{int(r['other_confirmed'])} |"
                    )
                lines.append("")
    (out / "manual_peak_overlap_details.md").write_text(
        "\n".join(lines) + "\n", encoding="utf-8"
    )
    print(f"written: {out}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", default=str(DOCX))
    ap.add_argument("--results", default="output/real_sets_cho_priority")
    ap.add_argument("--out", default="output/manual_peak_overlap_details")
    ap.add_argument("--datasets", default="all")
    args = ap.parse_args()
    _run(Path(args.docx), Path(args.results), Path(args.out), args.datasets)


if __name__ == "__main__":
    main()
