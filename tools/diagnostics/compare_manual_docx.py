"""Сопоставить ручные таблицы из DOCX с результатами пайплайна.

Ручные данные: 11 пиков F14 и 11 пиков F16, для 10 из них есть оценки
COOH/OH (пики >600 m/z исключены автором). Сравнение идёт по m/z и формуле;
дополнительно проверяется, входит ли ручная формула в all_candidates.
"""

from __future__ import annotations

import argparse
import ast
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document

DOCX = Path(
    "C:/Users/mvs/Work/Лаборатория/Аналитика/Тарабрин аналитика альфа-версия.docx"
)
DATASETS = {
    "Fulvagra F14": {"mass_table": 0, "count_table": 2, "prefix": "14"},
    "Fulvagra F16": {"mass_table": 1, "count_table": 3, "prefix": "16"},
}


def _clean(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split())


def _num(value: str) -> int | None:
    m = re.match(r"\s*(\d+)", _clean(value))
    return int(m.group(1)) if m else None


def _alt(value: str) -> int | None:
    m = re.search(r"\((\d+)\)", _clean(value))
    return int(m.group(1)) if m else None


def _manual_rows(doc: Document, cfg: dict) -> list[dict]:
    mass_rows = doc.tables[cfg["mass_table"]].rows
    count_rows = doc.tables[cfg["count_table"]].rows
    counts: dict[str, dict] = {}
    for row in count_rows[2:]:
        cells = [_clean(c.text) for c in row.cells]
        if len(cells) >= 4 and cells[0]:
            counts[cells[0]] = {
                "formula": cells[1],
                "cooh": _num(cells[2]),
                "cooh_alt": _alt(cells[2]),
                "oh": _num(cells[3]),
                "oh_alt": _alt(cells[3]),
            }
    out = []
    for row in mass_rows[2:]:
        cells = [_clean(c.text) for c in row.cells]
        if len(cells) < 3 or not cells[0]:
            continue
        num = cells[0]
        if num not in counts:
            continue
        item = {
            "number": num,
            "mass": float(cells[1]),
            "formula": cells[2],
            **counts[num],
        }
        out.append(item)
    return out


def _candidates(value) -> list[str]:
    if isinstance(value, list):
        return [str(v) for v in value]
    if isinstance(value, str) and value.strip():
        try:
            parsed = ast.literal_eval(value)
            return [str(v) for v in parsed] if isinstance(parsed, list) else []
        except (ValueError, SyntaxError):
            return []
    return []


def _nearest(df: pd.DataFrame, mass: float) -> pd.Series:
    idx = (df["mass"].astype(float) - mass).abs().idxmin()
    return df.loc[idx]


def _match(value: int | None, primary: int | None, alt: int | None) -> str:
    if value is None or primary is None:
        return "—"
    if value == primary:
        return "да"
    if alt is not None and value == alt:
        return "да (альт.)"
    return "нет"


def compare(
    docx: Path,
    result_root: Path,
    old_root: Path,
    out: Path,
    datasets: str = "all",
) -> None:
    doc = Document(docx)
    rows = []
    wanted = {x.strip().lower() for x in datasets.split(",")}
    selected = [
        name
        for name in DATASETS
        if datasets.strip().lower() in ("all", "real") or name.lower() in wanted
    ]
    for dataset in selected:
        cfg = DATASETS[dataset]
        new_path = result_root / dataset / "result_overlap_only.csv.gz"
        if not new_path.is_file():
            new_path = result_root / dataset / "result.csv.gz"
        new_df = pd.read_csv(new_path, low_memory=False)
        old_df = pd.read_csv(old_root / dataset / "result.csv.gz", low_memory=False)
        for manual in _manual_rows(doc, cfg):
            new = _nearest(new_df, manual["mass"])
            old = _nearest(old_df, manual["mass"])
            cands = _candidates(new.get("all_candidates"))
            rows.append(
                {
                    "dataset": dataset,
                    "number": manual["number"],
                    "manual_mass": manual["mass"],
                    "manual_formula": manual["formula"],
                    "manual_COOH": manual["cooh"],
                    "manual_COOH_alt": manual["cooh_alt"],
                    "manual_OH": manual["oh"],
                    "manual_OH_alt": manual["oh_alt"],
                    "software_mass": float(new["mass"]),
                    "delta_ppm": round(
                        (float(new["mass"]) - manual["mass"]) / manual["mass"] * 1e6,
                        4,
                    ),
                    "software_formula": str(new["brutto"]),
                    "formula_exact": str(new["brutto"]) == manual["formula"],
                    "formula_in_candidates": manual["formula"] in cands,
                    "software_COOH_raw": int(new["N_COOH_raw"]),
                    "software_OH_raw": int(new["N_OH_raw"]),
                    "software_COOH_new": int(new["N_COOH"]),
                    "software_OH_new": int(new["N_OH"]),
                    "software_COOH_old": int(old["N_COOH"]),
                    "software_OH_old": int(old["N_OH"]),
                    "COOH_raw_match": _match(
                        int(new["N_COOH_raw"]),
                        manual["cooh"],
                        manual["cooh_alt"],
                    ),
                    "OH_raw_match": _match(
                        int(new["N_OH_raw"]), manual["oh"], manual["oh_alt"]
                    ),
                    "COOH_new_match": _match(
                        int(new["N_COOH"]), manual["cooh"], manual["cooh_alt"]
                    ),
                    "OH_new_match": _match(
                        int(new["N_OH"]), manual["oh"], manual["oh_alt"]
                    ),
                }
            )

    df = pd.DataFrame(rows)
    out.mkdir(parents=True, exist_ok=True)
    df.to_csv(out / "manual_vs_pipeline.csv", index=False)
    summary = {
        "rows": len(df),
        "formula_exact": int(df["formula_exact"].sum()),
        "formula_in_candidates": int(df["formula_in_candidates"].sum()),
        "COOH_raw_match": int((df["COOH_raw_match"] != "нет").sum()),
        "OH_raw_match": int((df["OH_raw_match"] != "нет").sum()),
        "COOH_new_match": int((df["COOH_new_match"] != "нет").sum()),
        "OH_new_match": int((df["OH_new_match"] != "нет").sum()),
    }
    (out / "manual_vs_pipeline.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))

    lines = [
        "# Сопоставление с ручной обработкой Тарабрина",
        "",
        "Ручные значения: 10 соединений F14 и 10 F16, для которых в DOCX",
        "есть оценки COOH/OH. Пики 14-11 и 16-11 исключены автором (>600 m/z).",
        "",
        "| Сет | № | m/z | Формула вручную | Формула ПО | Δ ppm | "
        "COOH ручн. | COOH raw | COOH new | OH ручн. | OH raw | OH new |",
        "|---|---|---:|---|---|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for r in rows:

        def lab(v, alt):
            return f"{v} ({alt})" if alt is not None else str(v)

        lines.append(
            f"| {r['dataset']} | {r['number']} | {r['manual_mass']:.5f} | "
            f"{r['manual_formula']} | {r['software_formula']} | "
            f"{r['delta_ppm']:+.3f} | {lab(r['manual_COOH'], r['manual_COOH_alt'])} | "
            f"{r['software_COOH_raw']} | {r['software_COOH_new']} | "
            f"{lab(r['manual_OH'], r['manual_OH_alt'])} | "
            f"{r['software_OH_raw']} | {r['software_OH_new']} |"
        )
    lines += [
        "",
        "`new` — новое правило overlap-only; `raw` — длина серии до фильтра.",
        f"Формула ПО точно совпала в {summary['formula_exact']} из {summary['rows']} строк;",
        f"ручная формула была среди кандидатов в {summary['formula_in_candidates']} из {summary['rows']}.",
    ]
    (out / "manual_vs_pipeline.md").write_text(
        "\n".join(lines) + "\n", encoding="utf-8"
    )
    print(f"written: {out}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", default=str(DOCX))
    ap.add_argument("--results", default="output/real_sets_overlap_only")
    ap.add_argument("--old-results", default="output/real_sets_2026-10-22")
    ap.add_argument("--out", default="output/manual_vs_pipeline")
    ap.add_argument("--datasets", default="all")
    args = ap.parse_args()
    compare(
        Path(args.docx),
        Path(args.results),
        Path(args.old_results),
        Path(args.out),
        args.datasets,
    )


if __name__ == "__main__":
    main()
